"""
create_template.py
Creates P6_Import_Template.xlsx with formatted headers, type-coloured rows,
dropdown validation, a _README sheet, and realistic sample data pulled from
p6_reference.xml.

Requires p6_reference.xml in the same folder (your own P6 XML export).

  Row 1 = field names   (dark blue)
  Row 2 = field types   (colour-coded by type)
  Row 3+ = your data

Workflow:
  1. Export any project from P6: File -> Export -> Primavera P6 XML
     Rename the file to p6_reference.xml, place alongside this script.
  2. Run: py create_template.py
  3. Fill in your project data (row 3+) on each sheet.
  4. Run: py aurora_p6_converter.py P6_Import_Template.xlsx
     Produces P6_Import_pass1.xml / pass2.xml / pass3.xml.
  5. Import all three into P6 in order.
"""

import os, sys
import openpyxl
import openpyxl.utils
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
import xml.etree.ElementTree as ET

NS  = "http://xmlns.oracle.com/Primavera/P6Professional/V18.8/API/BusinessObjects"
XSI = "http://www.w3.org/2001/XMLSchema-instance"

# ── Row-1 style (dark blue, white bold) ──────────────────────────────────────
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
DATA_FONT = Font(name="Calibri", size=10)

# ── Row-2 fill by field type ──────────────────────────────────────────────────
_T = {                                                    # bg colour
    "Date":     ("70AD47", "FFFFFF"),   # green  / white text
    "String":   ("9DC3E6", "1F1F1F"),   # lt blue / dark text
    "Enum":     ("ED7D31", "FFFFFF"),   # orange / white text
    "Cost":     ("FFD966", "1F1F1F"),   # yellow / dark text
    "Unit":     ("FFD966", "1F1F1F"),
    "Duration": ("FFD966", "1F1F1F"),
    "Float":    ("FFD966", "1F1F1F"),
    "Integer":  ("FFD966", "1F1F1F"),
    "Lookup":   ("F4CCCD", "1F1F1F"),   # pink   / dark text
    "ObjectId": ("BFBFBF", "1F1F1F"),   # gray   / dark text
    "Boolean":  ("D9D2E9", "1F1F1F"),   # lavender / dark
}

def _type_fill(t):
    bg, fg = _T.get(t, ("2E75B6", "FFFFFF"))
    return PatternFill("solid", fgColor=bg), Font(italic=True, color=fg, name="Calibri", size=9)

# ── Dropdown allowed values ────────────────────────────────────────────────────
DROPDOWNS = {
    "Activity": {
        "Type": [
            "Task Dependent", "Resource Dependent", "Level of Effort",
            "WBS Summary", "Start Milestone", "Finish Milestone",
        ],
        "Status": ["Not Started", "In Progress", "Completed"],
        "PercentCompleteType": ["Physical", "Duration", "Units"],
        "DurationType": [
            "Fixed Duration and Units", "Fixed Duration and Units/Time",
            "Fixed Units", "Fixed Units/Time",
        ],
        "PrimaryConstraintType": [
            "Must Start On", "Must Finish By",
            "Start On or Before", "Start On or After",
            "Finish On or Before", "Finish On or After",
            "Mandatory Start", "Mandatory Finish",
        ],
    },
    "Relationship": {
        "Type": ["Finish to Start", "Start to Start", "Finish to Finish", "Start to Finish"],
    },
    "Expense": {
        "AccrualType": ["Uniform Over Activity", "Start of Activity", "End of Activity"],
    },
    "Resource": {
        "ResourceType": ["Labor", "NonLabor", "Material"],
    },
    "ResourceAssignment": {
        "ResourceType": ["Labor", "NonLabor", "Material"],
        "RateType": ["Price / Unit", "Unit / Time"],
    },
    "UDFValue": {
        "ObjectType": ["Activity", "WBS", "Project", "Resource"],
        "IndicatorValue": ["Red", "Yellow", "Green", "Blue"],
    },
    "WBS": {
        "Status": ["Active", "Inactive"],
    },
    "UDFType": {
        "DataType": ["Text", "Integer", "Double", "Cost", "Start Date", "Finish Date",
                     "Indicator", "Code"],
        "SubjectArea": ["Activity", "WBS", "Project", "Resource",
                        "Resource Assignment", "EPS", "Project Issue"],
    },
}

# ── Sheet helpers ─────────────────────────────────────────────────────────────
def style_hdr(ws):
    """Apply dark-blue style to row 1."""
    for cell in ws[1]:
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

def style_types(ws, types):
    """Apply per-type colour to row 2 cells."""
    for i, cell in enumerate(ws[2]):
        t = types[i] if i < len(types) else ""
        fill, font = _type_fill(t)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")

def style_data(ws):
    for r in range(3, ws.max_row + 1):
        for cell in ws[r]:
            cell.font = DATA_FONT

def auto_width(ws):
    for col in ws.columns:
        mx = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[col[0].column_letter].width = min(max(mx + 2, 12), 42)

def add_dropdowns(ws, headers, sheet_name):
    """Add DataValidation dropdowns to Enum columns."""
    dvs = DROPDOWNS.get(sheet_name, {})
    for col_name, values in dvs.items():
        if col_name not in headers:
            continue
        idx = headers.index(col_name) + 1
        col_letter = openpyxl.utils.get_column_letter(idx)
        formula = '"' + ",".join(values) + '"'
        dv = DataValidation(type="list", formula1=formula,
                            allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        dv.sqref = f"{col_letter}3:{col_letter}500"

def write_sheet(wb, name, headers, types, rows):
    ws = wb.create_sheet(name)
    ws.append(headers)
    ws.append(types)
    for row in rows:
        ws.append(row)
    style_hdr(ws)
    style_types(ws, types)
    style_data(ws)
    ws.freeze_panes = "A3"
    auto_width(ws)
    add_dropdowns(ws, headers, name)
    return ws

# ── README sheet ──────────────────────────────────────────────────────────────
def write_readme(wb):
    ws = wb.create_sheet("_README")
    ws.sheet_properties.tabColor = "00B050"
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 72

    SEC_FILL = PatternFill("solid", fgColor="1F4E79")
    SEC_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
    LBL_FONT = Font(bold=True, color="1F4E79", name="Calibri", size=10)
    TXT_FONT = Font(name="Calibri", size=10)
    TTL_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=14)
    TTL_FILL = PatternFill("solid", fgColor="1F4E79")
    SUB_FONT = Font(italic=True, color="FFFFFF", name="Calibri", size=10)

    def row(a="", b="", a_font=None, b_font=None, a_fill=None, b_fill=None, merge=False):
        ws.append([a, b])
        r = ws.max_row
        ca, cb = ws.cell(r, 1), ws.cell(r, 2)
        if a_font:  ca.font = a_font
        if b_font:  cb.font = b_font
        if a_fill:  ca.fill = a_fill
        if b_fill:  cb.fill = b_fill
        if merge:
            ws.merge_cells(f"A{r}:B{r}")
            ca.alignment = Alignment(horizontal="center", vertical="center")
        else:
            ca.alignment = Alignment(vertical="center", wrap_text=False)
            cb.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 18

    def section(title):
        row()
        ws.append([title, ""])
        r = ws.max_row
        for c in (ws.cell(r, 1), ws.cell(r, 2)):
            c.fill = SEC_FILL
            c.font = SEC_FONT
            c.alignment = Alignment(vertical="center")
        ws.row_dimensions[r].height = 18

    def legend_row(type_name, description):
        ws.append([type_name, description])
        r = ws.max_row
        bg, fg = _T.get(type_name, ("2E75B6", "FFFFFF"))
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=bg)
        ws.cell(r, 1).font = Font(italic=True, color=fg, name="Calibri", size=10, bold=True)
        ws.cell(r, 2).font = TXT_FONT
        ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 2).alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 18

    # ── Title ──────────────────────────────────────────────────────────────────
    ws.append(["AURORA P6 Excel Converter  —  v1.0", ""])
    ws.merge_cells("A1:B1")
    ws.cell(1, 1).font = TTL_FONT
    ws.cell(1, 1).fill = TTL_FILL
    ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.append(["Template for Primavera P6 Professional 18.8", ""])
    ws.merge_cells("A2:B2")
    ws.cell(2, 1).font = SUB_FONT
    ws.cell(2, 1).fill = PatternFill("solid", fgColor="2E75B6")
    ws.cell(2, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # ── How to Use ─────────────────────────────────────────────────────────────
    section("HOW TO USE")
    row("Step 1", 'Get your p6_reference.xml: In P6 Professional, go to File → Export → '
        'Primavera P6 XML. Export any project, rename the file to p6_reference.xml, '
        'and place it in the same folder as this template.',
        LBL_FONT, TXT_FONT)
    row("Step 2", "Generate a fresh template: run  py create_template.py  "
        "(or click 'Generate Template' in the GUI). "
        "This file is recreated with sample rows from your p6_reference.xml.",
        LBL_FONT, TXT_FONT)
    row("Step 3", "Fill in your project data on each sheet, starting from Row 3. "
        "Do NOT edit Row 1 (field names) or Row 2 (field types).",
        LBL_FONT, TXT_FONT)
    row("Step 4", "Run the converter:  py aurora_p6_converter.py P6_Import_Template.xlsx  "
        "(or click 'Convert to XML' in the GUI). "
        "Produces P6_Import_pass1.xml, pass2.xml, and pass3.xml next to this file.",
        LBL_FONT, TXT_FONT)
    row("Step 5", "Import all three XML files into P6 in order — see Import Instructions below.",
        LBL_FONT, TXT_FONT)

    # ── Import Instructions ────────────────────────────────────────────────────
    section("3-PASS IMPORT INSTRUCTIONS  (P6 Professional 18.8)")
    row("Pass 1", "File → Import → Primavera P6 XML → select P6_Import_pass1.xml "
        "→ Import action: Create New Project.  "
        "Imports: Project, WBS, Activities, Resources, ActivityCodes.",
        LBL_FONT, TXT_FONT)
    row("Pass 2", "File → Import → Primavera P6 XML → select P6_Import_pass2.xml "
        "→ Import action: Update Existing Project.  "
        "Imports: Activities (matched by ObjectId) + Relationships.",
        LBL_FONT, TXT_FONT)
    row("Pass 3", "File → Import → Primavera P6 XML → select P6_Import_pass3.xml "
        "→ Import action: Update Existing Project.  "
        "Imports: Activities (matched by ObjectId) + ResourceAssignments.  "
        "Pass 3 only required if ResourceAssignment sheet has data.",
        LBL_FONT, TXT_FONT)
    row("⚠ Note", "P6 18.8 raises FK constraint errors (fk_taskpred_task, fk_taskactv_task) "
        "if Relationships or ResourceAssignments are imported in the same transaction as "
        "new Activities. The 3-pass approach is required.",
        Font(bold=True, color="C00000", name="Calibri", size=10), TXT_FONT)

    # ── Sheet Descriptions ─────────────────────────────────────────────────────
    section("SHEET DESCRIPTIONS")
    SHEETS = [
        ("_README",            "This sheet — documentation and instructions (do not edit)."),
        ("_Config",            "P6 environment ObjectIds: CalendarObjectId, OBSObjectId, CurrencyObjectId. "
                               "Update these to match your P6 database before importing."),
        ("Project",            "One row per project. Fill: Id, Name, PlannedStartDate, "
                               "ScheduledFinishDate, ParentEPSObjectId."),
        ("WBS",                "Work Breakdown Structure hierarchy. Each row is one WBS node. "
                               "Use ParentObjectId to build the tree."),
        ("Activity",           "One row per schedule activity. Enum dropdowns provided for "
                               "Type, Status, PercentCompleteType, DurationType, PrimaryConstraintType."),
        ("Relationship",       "Activity dependencies. Each row links a predecessor to a successor "
                               "with a Type (FS/SS/FF/SF) and optional Lag."),
        ("ResourceAssignment", "Links a Resource to an Activity with planned/actual/remaining "
                               "units, costs, and dates."),
        ("UDFValue",           "User-defined field values. Set ObjectType to Activity, WBS, or Project; "
                               "ObjectId_Ref to the ActivityId or WBS Code; UDFTypeTitle to match a "
                               "title in the UDFType sheet."),
        ("UDFType",            "UDF type definitions — auto-populated from p6_reference.xml. "
                               "Do not change ObjectIds; they must match your P6 database."),
        ("Resource",           "Resource pool — auto-populated from p6_reference.xml. "
                               "Add new resources or update existing ones."),
        ("ActivityCodeType",   "Activity code type definitions (global or project-scoped)."),
        ("ActivityCode",       "Activity code values. Each row belongs to a CodeType via "
                               "CodeTypeObjectId."),
        ("Expense",            "Non-resource expense items attached to activities. "
                               "AccrualType controls when cost is recognised."),
    ]
    for name, desc in SHEETS:
        row(name, desc, LBL_FONT, TXT_FONT)

    # ── Field Type Legend ──────────────────────────────────────────────────────
    section("FIELD TYPE LEGEND  (Row 2 colour coding)")
    legend_row("Date",     "ISO 8601 datetime — e.g.  2024-03-15T08:00:00  or  2024-03-15")
    legend_row("String",   "Free text value — any characters allowed")
    legend_row("Enum",     "One of the allowed values — dropdown arrow appears in data cells")
    legend_row("Cost",     "Numeric cost value — e.g.  15000.00")
    legend_row("Unit",     "Resource units — e.g.  8.000000  (hours)")
    legend_row("Duration", "Duration in hours — e.g.  40.000000")
    legend_row("Float",    "Decimal number — e.g.  0.75")
    legend_row("Integer",  "Whole number — e.g.  10")
    legend_row("Lookup",   "Id or Code that references a row in another sheet")
    legend_row("ObjectId", "P6 database integer ID — must match your P6 environment")
    legend_row("Boolean",  "0 = false / No,   1 = true / Yes")

    # ── Important Notes ────────────────────────────────────────────────────────
    section("IMPORTANT NOTES")
    NOTES = [
        ("p6_reference.xml",
         "Required before running create_template.py. Export any existing project from "
         "P6 Professional (File → Export → Primavera P6 XML), rename to p6_reference.xml. "
         "Contains Calendar, OBS, Currency, UDFType, Role objects specific to your database."),
        ("Row 1 & Row 2",
         "Do NOT edit the header rows. Row 1 = field names used by the converter. "
         "Row 2 = field type reference (informational only, not parsed)."),
        ("ObjectIds in sample rows",
         "Sample ObjectIds are copied from p6_reference.xml and are valid for that database. "
         "When building a new project, assign new ObjectIds that don't conflict with existing data."),
        ("p6_reference.xml security",
         "This file contains real P6 database IDs — do not commit it to source control. "
         "It is listed in .gitignore."),
        ("Excel date format",
         "If Excel auto-formats Date cells, the converter still reads the underlying value "
         "correctly. For safety, enter dates as text:  2024-03-15T08:00:00"),
    ]
    for lbl, note in NOTES:
        row(f"• {lbl}", note, LBL_FONT, TXT_FONT)

    return ws



def main(work_dir=None):
    if work_dir is None:
        work_dir = os.path.dirname(os.path.abspath(__file__))
    # ── Load p6_reference.xml ────────────────────────────────────────────────────
    if not os.path.exists(os.path.join(work_dir, "p6_reference.xml")):
        print("ERROR: p6_reference.xml not found.")
        print("  Export any project from P6: File -> Export -> Primavera P6 XML")
        print("  Rename it to p6_reference.xml and place it in this folder.")
        sys.exit(1)

    tree = ET.parse(os.path.join(work_dir, "p6_reference.xml"))
    root = tree.getroot()
    proj = root.find(f"{{{NS}}}Project")

    def fv(el, tag):
        ch = el.find(f"{{{NS}}}{tag}")
        if ch is None:
            return None
        if ch.get(f"{{{XSI}}}nil", "false") == "true":
            return None
        return ch.text

    # ── UDFType lookup: ObjectId -> (Title, DataType) ────────────────────────────
    udf_oid_to_title   = {}
    udf_oid_to_dtype   = {}
    for u in root.findall(f"{{{NS}}}UDFType"):
        oid, title, dtype = fv(u, "ObjectId"), fv(u, "Title"), fv(u, "DataType")
        if oid and title:
            udf_oid_to_title[oid] = title
            udf_oid_to_dtype[oid] = dtype or ""

    # ── Expand sample set: up to 10 activities (more gives richer Rel/RA samples) ─
    ALL_ACTS       = proj.findall(f"{{{NS}}}Activity")
    sample_acts    = ALL_ACTS[:10]
    sample_act_oids = set(fv(a, "ObjectId") for a in sample_acts)
    oid_to_id       = {fv(a, "ObjectId"): fv(a, "Id") for a in sample_acts}

    # ── _CONFIG sheet ─────────────────────────────────────────────────────────────
    cfg_headers = ["Field", "Value", "Description"]
    cfg_types   = ["String", "ObjectId", "String"]
    cfg_rows = [
        ["CalendarObjectId", fv(proj, "ActivityDefaultCalendarObjectId"),
         "ObjectId of the global Calendar used by activities"],
        ["OBSObjectId",      fv(proj, "OBSObjectId"),
         "ObjectId of the OBS node assigned to this project and WBS"],
        ["CurrencyObjectId", "1",
         "ObjectId of the Currency (1 = USD Dollar in most databases)"],
    ]

    # ── PROJECT sheet ─────────────────────────────────────────────────────────────
    proj_headers = [
        "ObjectId","Id","Name","Status","OBSObjectId","ParentEPSObjectId",
        "WBSObjectId","PlannedStartDate","ScheduledFinishDate","DataDate",
        "ActivityIdPrefix","ActivityIdSuffix","ActivityIdIncrement",
        "ActivityDefaultCalendarObjectId","ActivityDefaultDurationType",
        "ActivityDefaultActivityType","ActivityDefaultPercentCompleteType",
        "OriginalBudget","LevelingPriority","GUID","AddedBy",
        "CriticalActivityPathType","CriticalActivityFloatLimit",
        "EarnedValueComputeType","EarnedValueETCComputeType",
        "WBSCodeSeparator","SummarizeToWBSLevel","SummaryLevel",
    ]
    proj_types = [
        "ObjectId","String","String","Enum","ObjectId","ObjectId",
        "ObjectId","Date","Date","Date",
        "String","String","Integer",
        "ObjectId","Enum","Enum","Enum",
        "Cost","Integer","String","String",
        "Enum","Duration","Enum","Enum",
        "String","Boolean","Enum",
    ]
    proj_data = [[
        fv(proj,"ObjectId"), fv(proj,"Id"), fv(proj,"Name"), fv(proj,"Status"),
        fv(proj,"OBSObjectId"), fv(proj,"ParentEPSObjectId"), fv(proj,"WBSObjectId"),
        fv(proj,"PlannedStartDate"), fv(proj,"ScheduledFinishDate"), fv(proj,"DataDate"),
        fv(proj,"ActivityIdPrefix"), fv(proj,"ActivityIdSuffix"), fv(proj,"ActivityIdIncrement"),
        fv(proj,"ActivityDefaultCalendarObjectId"), fv(proj,"ActivityDefaultDurationType"),
        fv(proj,"ActivityDefaultActivityType"), fv(proj,"ActivityDefaultPercentCompleteType"),
        fv(proj,"OriginalBudget"), fv(proj,"LevelingPriority"),
        fv(proj,"GUID"), fv(proj,"AddedBy"),
        fv(proj,"CriticalActivityPathType"), fv(proj,"CriticalActivityFloatLimit"),
        fv(proj,"EarnedValueComputeType"), fv(proj,"EarnedValueETCComputeType"),
        fv(proj,"WBSCodeSeparator"), fv(proj,"SummarizeToWBSLevel"), fv(proj,"SummaryLevel"),
    ]]

    # ── WBS sheet ─────────────────────────────────────────────────────────────────
    wbs_headers = [
        "ObjectId","Code","Name","ProjectObjectId","ParentObjectId",
        "OBSObjectId","Status","SequenceNumber","OriginalBudget",
        "EarnedValueComputeType","EarnedValueETCComputeType",
        "EarnedValueETCUserValue","EarnedValueUserPercent",
        "IndependentETCLaborUnits","IndependentETCTotalCost","GUID",
        "UDF_TypeObjectId","UDF_IndicatorValue",
    ]
    wbs_types = [
        "ObjectId","String","String","ObjectId","ObjectId",
        "ObjectId","Enum","Integer","Cost",
        "Enum","Enum","Float","Float",
        "Unit","Cost","String",
        "ObjectId","Enum",
    ]
    wbs_rows = []
    for w in proj.findall(f"{{{NS}}}WBS"):
        udf_el   = w.find(f"{{{NS}}}UDF")
        udf_type = fv(udf_el, "TypeObjectId")   if udf_el is not None else None
        udf_ind  = fv(udf_el, "IndicatorValue") if udf_el is not None else None
        wbs_rows.append([
            fv(w,"ObjectId"), fv(w,"Code"), fv(w,"Name"),
            fv(w,"ProjectObjectId"), fv(w,"ParentObjectId"),
            fv(w,"OBSObjectId"), fv(w,"Status"), fv(w,"SequenceNumber"),
            fv(w,"OriginalBudget"),
            fv(w,"EarnedValueComputeType"), fv(w,"EarnedValueETCComputeType"),
            fv(w,"EarnedValueETCUserValue"), fv(w,"EarnedValueUserPercent"),
            fv(w,"IndependentETCLaborUnits"), fv(w,"IndependentETCTotalCost"),
            fv(w,"GUID"), udf_type, udf_ind,
        ])

    # ── ACTIVITY sheet (up to 10 samples) ────────────────────────────────────────
    act_headers = [
        "ObjectId","Id","Name","ProjectObjectId","WBSObjectId",
        "Type","Status","CalendarObjectId",
        "PlannedStartDate","PlannedFinishDate","PlannedDuration",
        "ActualStartDate","ActualFinishDate","ActualDuration",
        "RemainingDuration","RemainingEarlyStartDate","RemainingEarlyFinishDate",
        "RemainingLateStartDate","RemainingLateFinishDate",
        "StartDate","FinishDate",
        "PercentComplete","PercentCompleteType","PhysicalPercentComplete",
        "DurationPercentComplete","UnitsPercentComplete",
        "PlannedLaborCost","PlannedLaborUnits","PlannedNonLaborCost","PlannedNonLaborUnits",
        "ActualLaborCost","ActualLaborUnits","ActualNonLaborCost","ActualNonLaborUnits",
        "RemainingLaborCost","RemainingLaborUnits","RemainingNonLaborCost","RemainingNonLaborUnits",
        "AtCompletionDuration","AtCompletionLaborCost","AtCompletionLaborUnits",
        "AtCompletionNonLaborCost","AtCompletionNonLaborUnits","AtCompletionExpenseCost",
        "DurationType","LevelingPriority","AutoComputeActuals",
        "PrimaryResourceObjectId","EstimatedWeight",
        "PrimaryConstraintType","PrimaryConstraintDate",
        "GUID",
        "Code1_TypeObjectId","Code1_ValueObjectId",
        "Code2_TypeObjectId","Code2_ValueObjectId",
        "Code3_TypeObjectId","Code3_ValueObjectId",
        "Code4_TypeObjectId","Code4_ValueObjectId",
    ]
    act_types = [
        "ObjectId","String","String","ObjectId","ObjectId",
        "Enum","Enum","ObjectId",
        "Date","Date","Duration",
        "Date","Date","Duration",
        "Duration","Date","Date",
        "Date","Date",
        "Date","Date",
        "Float","Enum","Float",
        "Float","Float",
        "Cost","Unit","Cost","Unit",
        "Cost","Unit","Cost","Unit",
        "Cost","Unit","Cost","Unit",
        "Duration","Cost","Unit",
        "Cost","Unit","Cost",
        "Enum","Enum","Boolean",
        "ObjectId","Float",
        "Enum","Date",
        "String",
        "ObjectId","ObjectId",
        "ObjectId","ObjectId",
        "ObjectId","ObjectId",
        "ObjectId","ObjectId",
    ]

    act_rows = []
    for a in sample_acts:
        codes = a.findall(f"{{{NS}}}Code")
        def gc(i, sub):
            if i < len(codes):
                ch = codes[i].find(f"{{{NS}}}{sub}")
                return ch.text if ch is not None else None
            return None
        act_rows.append([
            fv(a,"ObjectId"), fv(a,"Id"), fv(a,"Name"),
            fv(a,"ProjectObjectId"), fv(a,"WBSObjectId"),
            fv(a,"Type"), fv(a,"Status"), fv(a,"CalendarObjectId"),
            fv(a,"PlannedStartDate"), fv(a,"PlannedFinishDate"), fv(a,"PlannedDuration"),
            fv(a,"ActualStartDate"), fv(a,"ActualFinishDate"), fv(a,"ActualDuration"),
            fv(a,"RemainingDuration"),
            fv(a,"RemainingEarlyStartDate"), fv(a,"RemainingEarlyFinishDate"),
            fv(a,"RemainingLateStartDate"),  fv(a,"RemainingLateFinishDate"),
            fv(a,"StartDate"), fv(a,"FinishDate"),
            fv(a,"PercentComplete"), fv(a,"PercentCompleteType"), fv(a,"PhysicalPercentComplete"),
            fv(a,"DurationPercentComplete"), fv(a,"UnitsPercentComplete"),
            fv(a,"PlannedLaborCost"), fv(a,"PlannedLaborUnits"),
            fv(a,"PlannedNonLaborCost"), fv(a,"PlannedNonLaborUnits"),
            fv(a,"ActualLaborCost"), fv(a,"ActualLaborUnits"),
            fv(a,"ActualNonLaborCost"), fv(a,"ActualNonLaborUnits"),
            fv(a,"RemainingLaborCost"), fv(a,"RemainingLaborUnits"),
            fv(a,"RemainingNonLaborCost"), fv(a,"RemainingNonLaborUnits"),
            fv(a,"AtCompletionDuration"), fv(a,"AtCompletionLaborCost"), fv(a,"AtCompletionLaborUnits"),
            fv(a,"AtCompletionNonLaborCost"), fv(a,"AtCompletionNonLaborUnits"),
            fv(a,"AtCompletionExpenseCost"),
            fv(a,"DurationType"), fv(a,"LevelingPriority"), fv(a,"AutoComputeActuals"),
            fv(a,"PrimaryResourceObjectId"), fv(a,"EstimatedWeight"),
            fv(a,"PrimaryConstraintType"), fv(a,"PrimaryConstraintDate"),
            fv(a,"GUID"),
            gc(0,"TypeObjectId"), gc(0,"ValueObjectId"),
            gc(1,"TypeObjectId"), gc(1,"ValueObjectId"),
            gc(2,"TypeObjectId"), gc(2,"ValueObjectId"),
            gc(3,"TypeObjectId"), gc(3,"ValueObjectId"),
        ])

    # ── RELATIONSHIP sheet ────────────────────────────────────────────────────────
    rel_headers = [
        "ObjectId",
        "PredecessorActivityObjectId","PredecessorActivityId","PredecessorProjectObjectId",
        "SuccessorActivityObjectId","SuccessorActivityId","SuccessorProjectObjectId",
        "Type","Lag",
    ]
    rel_types = [
        "ObjectId",
        "ObjectId","String","ObjectId",
        "ObjectId","String","ObjectId",
        "Enum","Duration",
    ]
    rel_rows = []
    for r in proj.findall(f"{{{NS}}}Relationship"):
        pred_oid = fv(r, "PredecessorActivityObjectId")
        succ_oid = fv(r, "SuccessorActivityObjectId")
        if pred_oid in sample_act_oids and succ_oid in sample_act_oids:
            rel_rows.append([
                fv(r,"ObjectId"),
                pred_oid, oid_to_id.get(pred_oid, ""), fv(r,"PredecessorProjectObjectId"),
                succ_oid, oid_to_id.get(succ_oid, ""), fv(r,"SuccessorProjectObjectId"),
                fv(r,"Type"), fv(r,"Lag"),
            ])

    # ── RESOURCEASSIGNMENT sheet ──────────────────────────────────────────────────
    ra_headers = [
        "ProjectId","ActivityId","ResourceId","ResourceType","RateType",
        "PlannedUnits","PlannedCost","PlannedUnitsPerTime",
        "PlannedStartDate","PlannedFinishDate",
        "ActualUnits","ActualCost",
        "RemainingUnits","RemainingCost","RemainingDuration",
        "IsPrimaryResource","Proficiency",
        "ObjectId","ActivityObjectId","ResourceObjectId","ProjectObjectId","WBSObjectId",
    ]
    ra_types = [
        "Lookup","Lookup","Lookup","Enum","Enum",
        "Unit","Cost","Float",
        "Date","Date",
        "Unit","Cost",
        "Unit","Cost","Duration",
        "Boolean","Enum",
        "ObjectId","ObjectId","ObjectId","ObjectId","ObjectId",
    ]
    ra_sample_rows = []
    for ra in proj.findall(f"{{{NS}}}ResourceAssignment"):
        act_oid = fv(ra, "ActivityObjectId")
        if act_oid not in sample_act_oids:
            continue
        ra_sample_rows.append([
            fv(ra,"ProjectId"),
            fv(ra,"ActivityId"),        fv(ra,"ResourceId"),
            fv(ra,"ResourceType"),      fv(ra,"RateType"),
            fv(ra,"PlannedUnits"),      fv(ra,"PlannedCost"),
            fv(ra,"PlannedUnitsPerTime"),
            fv(ra,"PlannedStartDate"),  fv(ra,"PlannedFinishDate"),
            fv(ra,"ActualUnits"),       fv(ra,"ActualCost"),
            fv(ra,"RemainingUnits"),    fv(ra,"RemainingCost"),
            fv(ra,"RemainingDuration"),
            fv(ra,"IsPrimaryResource"), fv(ra,"Proficiency"),
            fv(ra,"ObjectId"),          act_oid,
            fv(ra,"ResourceObjectId"),  fv(ra,"ProjectObjectId"),
            fv(ra,"WBSObjectId"),
        ])
        if len(ra_sample_rows) >= 5:
            break

    # ── UDFTYPE sheet ─────────────────────────────────────────────────────────────
    udftype_headers = ["ObjectId","DataType","SubjectArea","Title","IsSecureCode"]
    udftype_types   = ["ObjectId","Enum","Enum","String","Boolean"]
    udftype_rows = []
    for u in root.findall(f"{{{NS}}}UDFType"):
        udftype_rows.append([
            fv(u,"ObjectId"), fv(u,"DataType"), fv(u,"SubjectArea"),
            fv(u,"Title"), fv(u,"IsSecureCode"),
        ])

    # ── UDFVALUE sheet ────────────────────────────────────────────────────────────
    udfv_headers = [
        "ProjectId","ObjectType","ObjectId_Ref","UDFTypeTitle",
        "TextValue","NumberValue","DateValue","IndicatorValue",
    ]
    udfv_types = [
        "Lookup","Enum","Lookup","Lookup",
        "String","Cost","Date","Enum",
    ]
    udfv_rows = []
    proj_id_ref = fv(proj, "Id")

    # Prefer Activity UDFs from reference (to avoid duplicating WBS sheet columns)
    for a in sample_acts:
        udf_el = a.find(f"{{{NS}}}UDF")
        if udf_el is None:
            continue
        type_oid = fv(udf_el, "TypeObjectId")
        title    = udf_oid_to_title.get(type_oid, "") if type_oid else ""
        if title:
            udfv_rows.append([
                proj_id_ref, "Activity", fv(a, "Id"), title,
                fv(udf_el, "TextValue"), fv(udf_el, "NumberValue"),
                fv(udf_el, "DateValue"), fv(udf_el, "IndicatorValue"),
            ])
        if len(udfv_rows) >= 3:
            break

    # Fallback: if reference has no Activity UDFs, create illustrative rows from UDFType list
    if not udfv_rows and udftype_rows and sample_acts:
        # Pick up to 3 UDFTypes with SubjectArea matching Activity (or any if none)
        act_types_udf = [r for r in udftype_rows if r[2] == "Activity"] or udftype_rows
        for i, udf_row in enumerate(act_types_udf[:3]):
            oid, dtype, _, title, _ = udf_row
            act = sample_acts[min(i, len(sample_acts)-1)]
            dtype = udf_oid_to_dtype.get(oid, dtype or "Text")
            vals = {
                "Text":      [f"Example text value {i+1}", None, None, None],
                "Indicator": [None, None, None, ["Green","Yellow","Red"][i % 3]],
                "Integer":   [None, str(10 * (i+1)), None, None],
                "Double":    [None, str(float(10 * (i+1))), None, None],
                "Cost":      [None, str(1000 * (i+1)) + ".00", None, None],
            }.get(dtype, [f"Example {i+1}", None, None, None])
            if title:
                udfv_rows.append([proj_id_ref, "Activity", fv(act, "Id"), title] + vals)

    # ── ACTIVITYCODETYPE sheet ────────────────────────────────────────────────────
    actype_headers = [
        "ObjectId","Name","Scope","Length","IsSecureCode","SequenceNumber","RefProjectObjectIds",
    ]
    actype_types = ["ObjectId","String","Enum","Integer","Boolean","Integer","String"]
    actype_rows = []
    for act in root.findall(f"{{{NS}}}ActivityCodeType"):
        actype_rows.append([
            fv(act,"ObjectId"), fv(act,"Name"), fv(act,"Scope"),
            fv(act,"Length"), fv(act,"IsSecureCode"), fv(act,"SequenceNumber"),
            fv(act,"RefProjectObjectIds"),
        ])

    # ── ACTIVITYCODE sheet ────────────────────────────────────────────────────────
    ac_headers = ["ObjectId","CodeTypeObjectId","CodeValue","Description","Color","SequenceNumber"]
    ac_types   = ["ObjectId","ObjectId","String","String","String","Integer"]
    ac_rows = []
    for ac in root.findall(f"{{{NS}}}ActivityCode"):
        ac_rows.append([
            fv(ac,"ObjectId"), fv(ac,"CodeTypeObjectId"), fv(ac,"CodeValue"),
            fv(ac,"Description"), fv(ac,"Color"), fv(ac,"SequenceNumber"),
        ])

    # ── RESOURCE sheet ────────────────────────────────────────────────────────────
    res_headers = [
        "ObjectId","Id","Name","Code","ResourceType","CalendarObjectId",
        "CurrencyObjectId","DefaultUnitsPerTime","OvertimeFactor",
        "IsActive","IsOverTimeAllowed","AutoComputeActuals",
        "CalculateCostFromUnits","SequenceNumber","GUID",
        "ParentObjectId","PrimaryRoleObjectId","ResourceNotes","UnitOfMeasureObjectId",
    ]
    res_types = [
        "ObjectId","String","String","String","Enum","ObjectId",
        "ObjectId","Unit","Float",
        "Boolean","Boolean","Boolean",
        "Boolean","Integer","String",
        "ObjectId","ObjectId","String","ObjectId",
    ]
    res_rows = []
    for r in root.findall(f"{{{NS}}}Resource"):
        res_rows.append([
            fv(r,"ObjectId"), fv(r,"Id"), fv(r,"Name"), fv(r,"Code"),
            fv(r,"ResourceType"), fv(r,"CalendarObjectId"),
            fv(r,"CurrencyObjectId"), fv(r,"DefaultUnitsPerTime"), fv(r,"OvertimeFactor"),
            fv(r,"IsActive"), fv(r,"IsOverTimeAllowed"), fv(r,"AutoComputeActuals"),
            fv(r,"CalculateCostFromUnits"), fv(r,"SequenceNumber"), fv(r,"GUID"),
            fv(r,"ParentObjectId"), fv(r,"PrimaryRoleObjectId"),
            fv(r,"ResourceNotes"), fv(r,"UnitOfMeasureObjectId"),
        ])

    # ── EXPENSE sheet ─────────────────────────────────────────────────────────────
    exp_headers = [
        "ObjectId","ActivityObjectId","ProjectObjectId","Description",
        "CostAccountObjectId","PlannedCost","ActualCost","RemainingCost",
        "AtCompletionCost","AutoComputeActuals","AccrualType","VendorId",
    ]
    exp_types = [
        "ObjectId","ObjectId","ObjectId","String",
        "ObjectId","Cost","Cost","Cost",
        "Cost","Boolean","Enum","String",
    ]
    # Sample Expense rows — 3 rows from the first 3 sample activities
    _expense_desc  = ["Survey & Geotechnical", "Equipment Rental", "Subcontractor Materials"]
    _expense_costs = ["5000.00", "12500.00", "8750.00"]
    exp_rows = []
    for i, a in enumerate(sample_acts[:3]):
        act_oid  = fv(a, "ObjectId")
        proj_oid = fv(a, "ProjectObjectId") or fv(proj, "ObjectId")
        exp_rows.append([
            None, act_oid, proj_oid,
            _expense_desc[i],
            None,                # CostAccountObjectId
            _expense_costs[i],   # PlannedCost
            "0.00",              # ActualCost
            _expense_costs[i],   # RemainingCost
            _expense_costs[i],   # AtCompletionCost
            "0",                 # AutoComputeActuals
            "Uniform Over Activity",  # AccrualType
            None,                # VendorId
        ])

    # ── Write workbook ─────────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    write_readme(wb)                                                         # first sheet
    write_sheet(wb, "_Config",            cfg_headers,     cfg_types,     cfg_rows)
    write_sheet(wb, "Project",            proj_headers,    proj_types,    proj_data)
    write_sheet(wb, "WBS",                wbs_headers,     wbs_types,     wbs_rows)
    write_sheet(wb, "Activity",           act_headers,     act_types,     act_rows)
    write_sheet(wb, "Relationship",       rel_headers,     rel_types,     rel_rows)
    write_sheet(wb, "ResourceAssignment", ra_headers,      ra_types,      ra_sample_rows)
    write_sheet(wb, "UDFType",            udftype_headers, udftype_types, udftype_rows)
    write_sheet(wb, "UDFValue",           udfv_headers,    udfv_types,    udfv_rows)
    write_sheet(wb, "ActivityCodeType",   actype_headers,  actype_types,  actype_rows)
    write_sheet(wb, "ActivityCode",       ac_headers,      ac_types,      ac_rows)
    write_sheet(wb, "Expense",            exp_headers,     exp_types,     exp_rows)
    write_sheet(wb, "Resource",           res_headers,     res_types,     res_rows)

    out = os.path.join(work_dir, "P6_Import_Template.xlsx")
    wb.save(out)
    print(f"Saved {out}")
    print(f"  _Config              : {len(cfg_rows)} rows")
    print(f"  Project              : {len(proj_data)} rows")
    print(f"  WBS                  : {len(wbs_rows)} rows")
    print(f"  Activity             : {len(act_rows)} rows  (up to 10 from p6_reference.xml)")
    print(f"  Relationship         : {len(rel_rows)} rows")
    print(f"  ResourceAssignment   : {len(ra_sample_rows)} rows")
    print(f"  UDFType              : {len(udftype_rows)} rows")
    print(f"  UDFValue             : {len(udfv_rows)} rows")
    print(f"  ActivityCodeType     : {len(actype_rows)} rows")
    print(f"  ActivityCode         : {len(ac_rows)} rows")
    print(f"  Expense              : {len(exp_rows)} rows  (3 sample expense items)")
    print(f"  Resource             : {len(res_rows)} rows")
    print(f"  Sheets with dropdowns: Activity, Relationship, Expense, Resource,")
    print(f"                         ResourceAssignment, UDFValue, UDFType, WBS")


if __name__ == '__main__':
    main()
